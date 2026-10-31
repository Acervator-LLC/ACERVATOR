"""Slide primitives, and the New User Procedure deck ``build_procedure`` composes.

``blank``, ``text_box``, ``picture``, ``heading``, ``source_note`` and ``table``
place one element on a slide; ``tools/build_product_manual.py`` composes the
evidence deck with them.

    python -m tools.build_vwap_deck --deck new-user
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING

import pptx
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Length, Pt

if TYPE_CHECKING:
    from pptx.presentation import Presentation
    from pptx.slide import Slide

REPO_ROOT = Path(__file__).resolve().parent.parent


BG = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTE = RGBColor(0x6B, 0x6B, 0x6B)
ACCENT = RGBColor(0x1F, 0x3A, 0x68)


def emu(value: float) -> Length:
    """Return one length in English metric units."""
    return Emu(int(value))


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.55)
BODY_W = emu(SLIDE_W - 2 * MARGIN)


class BuildRefusedError(RuntimeError):
    """The deck cannot be composed, and the message names what is missing."""


@dataclass(frozen=True)
class Box:
    """One rectangle on a slide, in English metric units."""

    left: Length
    top: Length
    width: Length
    height: Length


@dataclass
class Ledger:
    """Every figure a build placed, so the embedded count is measured."""

    placed: list[Path] = field(default_factory=list)
    stated_absent: list[str] = field(default_factory=list)

    def embedded(self) -> int:
        """Return how many distinct figure files this build placed."""
        return len({path.name for path in self.placed})


def tables(text: str) -> list[list[list[str]]]:
    """Return every pipe table in the page, each one a list of cell rows."""
    found: list[list[list[str]]] = []
    rows: list[list[str]] = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if not all(cell and set(cell) <= set("-: ") for cell in cells):
                rows.append(cells)
            continue
        if rows:
            found.append(rows)
            rows = []
    if rows:
        found.append(rows)
    return found


def table_with_header(text: str, header: list[str]) -> list[dict[str, str]]:
    """Return the rows of the one table whose header row is ``header``."""
    for block in tables(text):
        if block and block[0] == header:
            return [dict(zip(header, row, strict=True)) for row in block[1:]]
    message = f"no table in the page carries the header {header}"
    raise BuildRefusedError(message)


def blank(prs: Presentation) -> Slide:
    """Add one slide on a white ground and return it."""
    slide: Slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    return slide


def text_box(
    slide: Slide,
    box: Box,
    body: str,
    *,
    size: int = 12,
    bold: bool = False,
    color: RGBColor = INK,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    """Place one wrapped text box, one paragraph per line of ``body``."""
    shape = slide.shapes.add_textbox(box.left, box.top, box.width, box.height)
    frame = shape.text_frame
    frame.word_wrap = True
    for index, line in enumerate(body.split("\n")):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def picture(slide: Slide, path: Path, box: Box, ledger: Ledger) -> None:
    """Place one figure inside the box, keeping its own aspect ratio."""
    shape = slide.shapes.add_picture(str(path), box.left, box.top, width=box.width)
    if shape.height > box.height:
        shape.width = emu(shape.width * box.height / shape.height)
        shape.height = box.height
    shape.left = emu(box.left + (box.width - shape.width) / 2)
    shape.top = emu(box.top + (box.height - shape.height) / 2)
    ledger.placed.append(path)


def heading(slide: Slide, title: str, reading: str = "") -> None:
    """Place a slide's title, and the one line that reads its figure."""
    text_box(
        slide,
        Box(MARGIN, Inches(0.34), BODY_W, Inches(0.7)),
        title,
        size=26,
        bold=True,
        color=ACCENT,
    )
    if reading:
        text_box(
            slide,
            Box(MARGIN, Inches(1.06), BODY_W, Inches(0.62)),
            reading,
            size=13,
            color=MUTE,
        )


def source_note(slide: Slide, note: str) -> None:
    """Place the line naming where a slide's content came from."""
    text_box(
        slide,
        Box(MARGIN, emu(SLIDE_H - Inches(0.52)), BODY_W, Inches(0.34)),
        note,
        size=9,
        color=MUTE,
        align=PP_ALIGN.RIGHT,
    )


def table(slide: Slide, header: list[str], rows: list[list[str]], box: Box) -> None:
    """Place one table, its header on the accent ground."""
    shape = slide.shapes.add_table(
        len(rows) + 1,
        len(header),
        box.left,
        box.top,
        box.width,
        box.height,
    )
    grid = shape.table
    for column, name in enumerate(header):
        cell = grid.cell(0, column)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
        cell.text = name
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = BG
    for index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            cell = grid.cell(index, column)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG
            cell.text = value
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(11)
            run.font.color.rgb = INK


PROCEDURE_DIR = REPO_ROOT / "docs" / "new-user-procedure"
PROCEDURE_README = PROCEDURE_DIR / "README.md"
PROCEDURE_DECK = PROCEDURE_DIR / "New-User-Procedure.pptx"
PROCEDURE_ORDER_HEADER = ["#", "page", "the step"]
PROCEDURE_PAGE_LINK = re.compile(r"\]\((\d\d-[^)]+\.md)\)")
PROCEDURE_HEADING = re.compile(r"^#[ \t]+(.+?)[ \t]*$", re.MULTILINE)
PROCEDURE_FIGURE = re.compile(r"!\[(.+?)\]\((figures/[^)]+)\)", re.DOTALL)
PROCEDURE_GUIDE = re.compile(r"^\*\*(.+?)\*\*[ \t]*$", re.MULTILINE | re.DOTALL)
PROCEDURE_PICTURE_BOX = Box(MARGIN, Inches(1.8), BODY_W, Inches(4.2))
PROCEDURE_GUIDE_BOX = Box(MARGIN, Inches(6.14), BODY_W, Inches(0.72))
PROCEDURE_GUIDE_SIZE = 15


@dataclass(frozen=True)
class ProcedureStep:
    """One step page's heading, picture caption, guiding line and figure."""

    page: str
    title: str
    caption: str
    guide: str
    figure: Path

    def note(self) -> str:
        """Return the line naming the page this step was read from."""
        return f"docs/new-user-procedure/{self.page}"


def fold_lines(text: str) -> str:
    """Return ``text`` on one line, every run of whitespace a single space."""
    return re.sub(r"\s+", " ", text).strip()


def procedure_page_names() -> list[str]:
    """Return the step page names in the order the README table sets."""
    if not PROCEDURE_README.is_file():
        message = f"the order table is not on disk: {PROCEDURE_README}"
        raise BuildRefusedError(message)
    rows = table_with_header(
        PROCEDURE_README.read_text(encoding="utf-8"),
        PROCEDURE_ORDER_HEADER,
    )
    names: list[str] = []
    for row in rows:
        link = PROCEDURE_PAGE_LINK.search(row["page"])
        if link is None:
            message = f"an order row links to no step page: {row}"
            raise BuildRefusedError(message)
        names.append(link.group(1))
    return names


def procedure_step(name: str) -> ProcedureStep:
    """Return one ``ProcedureStep`` read from its page, refusing an absent part."""
    path = PROCEDURE_DIR / name
    if not path.is_file():
        message = f"the order table names a page the disk does not carry: {path}"
        raise BuildRefusedError(message)
    text = path.read_text(encoding="utf-8")
    title = PROCEDURE_HEADING.search(text)
    if title is None:
        message = f"{name} carries no heading"
        raise BuildRefusedError(message)
    drawn = PROCEDURE_FIGURE.search(text)
    if drawn is None:
        message = f"{name} embeds no picture"
        raise BuildRefusedError(message)
    guide = PROCEDURE_GUIDE.search(text)
    if guide is None:
        message = f"{name} carries no guiding line"
        raise BuildRefusedError(message)
    figure = PROCEDURE_DIR / drawn.group(2)
    if not figure.is_file():
        message = f"{name} embeds a figure the disk does not carry: {figure}"
        raise BuildRefusedError(message)
    return ProcedureStep(
        page=name,
        title=fold_lines(title.group(1)),
        caption=fold_lines(drawn.group(1)),
        guide=fold_lines(guide.group(1)),
        figure=figure,
    )


def slide_procedure_step(
    prs: Presentation,
    step: ProcedureStep,
    ledger: Ledger,
) -> None:
    """Add one slide carrying the step's heading, caption, picture and guide."""
    slide = blank(prs)
    heading(slide, step.title, step.caption)
    picture(slide, step.figure, PROCEDURE_PICTURE_BOX, ledger)
    text_box(
        slide, PROCEDURE_GUIDE_BOX, step.guide, size=PROCEDURE_GUIDE_SIZE, bold=True
    )
    source_note(slide, step.note())


def compose_procedure(steps: list[ProcedureStep]) -> tuple[Presentation, Ledger]:
    """Return the procedure deck, one slide per step, and the figure ledger."""
    ledger = Ledger()
    prs = pptx.Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for step in steps:
        slide_procedure_step(prs, step, ledger)
    return prs, ledger


def slide_lines(slide: Slide) -> list[str]:
    """Return every non-empty line of text the slide draws, in shape order."""
    found: list[str] = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            line = fold_lines("".join(run.text for run in para.runs))
            if line:
                found.append(line)
    return found


def slide_images(slide: Slide) -> list[bytes]:
    """Return the bytes of every picture the slide embeds, in shape order."""
    found: list[bytes] = []
    for shape in slide.shapes:
        image = getattr(shape, "image", None)
        if image is not None:
            found.append(image.blob)
    return found


def verify_procedure_deck(path: Path, steps: list[ProcedureStep]) -> None:
    """Read the written deck back and refuse a slide that differs from its page."""
    written = pptx.Presentation(str(path))
    slides = list(written.slides)
    if len(slides) != len(steps):
        message = (
            f"the written deck holds {len(slides)} slides"
            f" against {len(steps)} pages"
        )
        raise BuildRefusedError(message)
    for index, (slide, step) in enumerate(zip(slides, steps, strict=True), start=1):
        lines = slide_lines(slide)
        absent = [
            wanted
            for wanted in (step.title, step.caption, step.guide, step.note())
            if wanted not in lines
        ]
        images = slide_images(slide)
        if absent or images != [step.figure.read_bytes()]:
            message = (
                f"slide {index} differs from {step.page}:"
                f" absent {absent}, pictures {len(images)}"
            )
            raise BuildRefusedError(message)
        print(f"slide {index:2d}  {step.page:34s}  title caption guide picture")


def build_procedure() -> Path:
    """Compose the procedure deck from its pages and write it beside them."""
    steps = [procedure_step(name) for name in procedure_page_names()]
    prs, ledger = compose_procedure(steps)
    prs.save(str(PROCEDURE_DECK))
    verify_procedure_deck(PROCEDURE_DECK, steps)
    size = PROCEDURE_DECK.stat().st_size
    print(f"slides            {len(prs.slides)}")
    print(f"figures embedded  {ledger.embedded()}")
    print(f"wrote             {PROCEDURE_DECK}  {size} bytes")
    return PROCEDURE_DECK


def build_parser() -> argparse.ArgumentParser:
    """Return the command line, whose ``--deck`` names the deck to compose."""
    parser = argparse.ArgumentParser(
        description="Compose the New User Procedure slide show.",
    )
    parser.add_argument(
        "--deck",
        choices=("new-user",),
        default="new-user",
        help="the deck to compose; new-user reads docs/new-user-procedure",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    """Compose the ``--deck`` deck, and report what was written."""
    build_parser().parse_args(argv)
    try:
        build_procedure()
    except BuildRefusedError as refusal:
        print(f"REFUSED           {refusal}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
